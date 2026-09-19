"""
DOCX 解析服务：基于 python-docx

输出格式与 PDFParser 完全一致（meta, page_count, full_text, toc, pages）
方便 chapter_splitter 等下游模块无需区分文件类型。

注意：DOCX 没有"页"概念，按段落分块（每 N 段一页），让下游逻辑一致。
"""
import re
from pathlib import Path
from typing import List, Dict, Any, Optional
from dataclasses import dataclass

from app.utils.logger import logger


@dataclass
class BBox:
    """占位（DOCX 没有真实 bbox）"""
    x0: float = 0
    y0: float = 0
    x1: float = 0
    y1: float = 0


@dataclass
class PageContent:
    page_num: int
    text: str
    width: float = 595.0
    height: float = 842.0
    blocks: List[Dict[str, Any]] = None

    def __post_init__(self):
        if self.blocks is None:
            self.blocks = []


@dataclass
class PaperMeta:
    title: str = ""
    authors: List[str] = None
    abstract: str = ""
    keywords: List[str] = None
    year: Optional[int] = None
    venue: Optional[str] = None
    doi: Optional[str] = None

    def __post_init__(self):
        if self.authors is None:
            self.authors = []
        if self.keywords is None:
            self.keywords = []


# 每多少段算一页（让 docx 的"页"对标 PDF 的视觉页）
PARAS_PER_PAGE = 25


class DOCXParser:
    """DOCX 解析器（接口与 PDFParser 一致）"""

    def __init__(self, docx_path: str | Path):
        try:
            from docx import Document
        except ImportError as e:
            raise ImportError(
                "缺少 python-docx 依赖，请运行: pip install python-docx"
            ) from e

        self.docx_path = Path(docx_path)
        if not self.docx_path.exists():
            raise FileNotFoundError(f"DOCX 不存在: {self.docx_path}")

        self.doc = Document(str(self.docx_path))
        self.paragraphs = self.doc.paragraphs
        self.total_paras = len(self.paragraphs)
        logger.info(f"打开 DOCX: {self.docx_path.name} ({self.total_paras} 段)")

    def close(self):
        """python-docx 不需要显式关闭"""
        pass

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc_val, exc_tb):
        self.close()

    # ============ 元数据提取 ============

    def extract_metadata(self) -> PaperMeta:
        """从 DOCX 提取标题、作者、摘要、关键词、年份"""
        if self.total_paras == 0:
            return PaperMeta()

        # 标题：第一个 Heading 1 / Title 样式，或第一段
        title = self._extract_title()
        # 作者：标题后到 Abstract 之间的内容
        authors = self._extract_authors(title)
        # 摘要：找 "Abstract" 段后的内容
        abstract = self._extract_abstract()
        # 关键词：找 "Keywords" 段
        keywords = self._extract_keywords()
        # 年份：全文搜索
        year = self._extract_year()

        return PaperMeta(
            title=title,
            authors=authors,
            abstract=abstract,
            keywords=keywords,
            year=year,
        )

    def _extract_title(self) -> str:
        """提取标题"""
        # 优先级 1: 第一个 Heading 1 / Title 样式
        for p in self.paragraphs:
            style_name = (p.style.name or "").strip()
            if style_name in ("Title", "Heading 1"):
                text = p.text.strip()
                if text and len(text) > 5:
                    return text
        # 优先级 2: 第一段非空文本
        for p in self.paragraphs:
            text = p.text.strip()
            if text and len(text) > 5:
                return text
        return ""

    def _extract_authors(self, title: str) -> List[str]:
        """提取作者"""
        if not title:
            return []
        # 标题后到 Abstract 之间的内容
        capture = False
        for p in self.paragraphs:
            text = p.text.strip()
            if not capture:
                # style.name 判空与同文件其他位置保持一致：样式损坏/缺失的 docx 会是 None
                if title in text or ((p.style.name or "").strip() in ("Title", "Heading 1")):
                    capture = True
                continue
            if not text:
                continue
            lower = text.lower()
            if "abstract" in lower or "摘 要" in lower or "摘要" in text:
                break
            # 包含逗号 / and / · 才像作者列表
            if "," in text or " and " in lower or "·" in text or "，" in text:
                # 拆分
                parts = re.split(r'[,，·;；]| and ', text)
                authors = []
                for part in parts:
                    part = part.strip().rstrip(".*†")
                    if part and 2 < len(part) < 50 and not part.isdigit():
                        authors.append(part)
                if authors:
                    return authors[:10]
        return []

    def _extract_abstract(self) -> str:
        """提取摘要"""
        capture = False
        abstract_paras = []
        for p in self.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            lower = text.lower()
            # 找 "Abstract" 开头（B11: 不再限制段落长度 —— 长段落摘要 "Abstract: ..." 不丢）
            if not capture:
                if lower.startswith("abstract") or "摘 要" in text or text.startswith("摘要"):
                    # 去掉 "Abstract"/"Abstract:" 前缀后作为首段摘要（可能就是完整摘要段）
                    cleaned = re.sub(r'^(abstract|摘\s*要|摘要)[:：\s]*', '', text, flags=re.IGNORECASE).strip()
                    if cleaned:
                        abstract_paras.append(cleaned)
                    capture = True
                continue
            # 已经在 abstract 区，遇到下一个章节标题就停
            style_name = (p.style.name or "").strip()
            if style_name.startswith("Heading") and style_name != "Heading 1":
                break
            if lower.startswith("introduction") or "1. introduction" in lower or "1 引言" in text:
                break
            abstract_paras.append(text)
            if sum(len(p) for p in abstract_paras) > 2000:
                break
        abstract = " ".join(abstract_paras)
        return re.sub(r'\s+', ' ', abstract)[:2000]

    def _extract_keywords(self) -> List[str]:
        """提取关键词"""
        capture = False
        for p in self.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            lower = text.lower()
            if not capture:
                if (lower.startswith("keyword") or "关键词" in text) and len(text) < 200:
                    capture = True
                    # 同段可能就有关键词
                    cleaned = re.sub(r'^(keywords?|关键词)[:：\s]*', '', text, flags=re.IGNORECASE).strip()
                    if cleaned:
                        kws = re.split(r'[,，;；·\n]+', cleaned)
                        return [k.strip() for k in kws if k.strip() and len(k.strip()) < 50][:10]
                continue
            # 已经进入 keywords 区，提取这一段
            kws = re.split(r'[,，;；·\n]+', text)
            result = [k.strip() for k in kws if k.strip() and len(k.strip()) < 50][:10]
            if result:
                return result
        return []

    def _extract_year(self) -> Optional[int]:
        """提取年份"""
        from collections import Counter
        # 全文搜索 1900-2099
        full_text = self.extract_full_text()
        matches = re.findall(r'\b(19\d{2}|20\d{2})\b', full_text[:10000])
        if not matches:
            return None
        return int(Counter(matches).most_common(1)[0][0])

    # ============ 文本提取 ============

    def extract_pages(self) -> List[PageContent]:
        """按段落分块成"伪页"（每 25 段一页）"""
        # 过滤空段
        non_empty = [p.text for p in self.paragraphs if p.text.strip()]
        if not non_empty:
            return [PageContent(page_num=1, text="")]

        pages = []
        for i in range(0, len(non_empty), PARAS_PER_PAGE):
            chunk = non_empty[i:i + PARAS_PER_PAGE]
            pages.append(PageContent(
                page_num=i // PARAS_PER_PAGE + 1,
                text="\n".join(chunk),
                width=595.0,
                height=842.0,
                blocks=[],
            ))
        return pages

    def extract_full_text(self) -> str:
        """提取全文（按"伪页"拼接：每 25 段一页，带 === Page N === 标记，与 PDF 版一致）

        页码标记保证 legacy 整篇分析（BREAKDOWN_SYSTEM_PROMPT 声明正文含标记）时
        LLM 的 quotes.page 有锚点可依，不凭空猜测。
        """
        non_empty = [p.text for p in self.paragraphs if p.text.strip()]
        if not non_empty:
            return ""
        texts = []
        for i in range(0, len(non_empty), PARAS_PER_PAGE):
            chunk = non_empty[i:i + PARAS_PER_PAGE]
            texts.append(f"\n\n=== Page {i // PARAS_PER_PAGE + 1} ===\n\n" + "\n".join(chunk))
        return "\n".join(texts)

    def extract_toc(self) -> List[List]:
        """提取大纲（基于 Heading 1 样式 + 估算页码）"""
        toc = []
        for i, p in enumerate(self.paragraphs):
            style_name = (p.style.name or "").strip()
            if style_name.startswith("Heading 1"):
                # 估算页码（按段分页）
                page = i // PARAS_PER_PAGE + 1
                title = p.text.strip()
                if title:
                    toc.append([1, title, page])
        return toc

    # ============ 一站式解析 ============

    def parse(self) -> Dict[str, Any]:
        """完整解析：返回与 PDFParser 同样格式的 dict"""
        meta = self.extract_metadata()
        pages = self.extract_pages()
        full_text = self.extract_full_text()
        toc = self.extract_toc()

        return {
            "meta": {
                "title": meta.title,
                "authors": meta.authors,
                "abstract": meta.abstract,
                "keywords": meta.keywords,
                "year": meta.year,
                "venue": meta.venue,
                "doi": meta.doi,
            },
            "page_count": len(pages),
            "full_text": full_text,
            "toc": toc,
            "pages": [
                {
                    "page_num": p.page_num,
                    "text": p.text,
                    "width": p.width,
                    "height": p.height,
                    "blocks": p.blocks or [],
                }
                for p in pages
            ],
        }


def parse_docx(docx_path: str | Path) -> Dict[str, Any]:
    """便捷函数：解析 DOCX"""
    with DOCXParser(docx_path) as parser:
        return parser.parse()
