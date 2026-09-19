"""
PaperPilot E2E 测试 - 基类 + helper + lifecycle（Sprint 4 / R2 拆分）
"""
import io
import os
import sys
import shutil
import time
from pathlib import Path
from contextlib import ExitStack
from unittest.mock import AsyncMock, patch

# Windows GBK 终端兼容
try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

# pytest 模式下 env 由 conftest.py 统一注入（先于本模块 import 生效，见 conftest.py），
# 这里仅兜底脚本模式（python -m tests.test_e2e，不经过 conftest）。
# setdefault 不会覆盖 conftest 已注入的值。
os.environ.setdefault("APP_API_KEY", "e2e-test-api-key-2026")
os.environ.setdefault("RATE_LIMIT_GENERAL_PER_MIN", "200")
os.environ.setdefault("RATE_LIMIT_EXPENSIVE_PER_MIN", "10")
os.environ.setdefault("LLM_API_KEY", "")             # 兼容老 DEEPSEEK_API_KEY 别名
os.environ.setdefault("TAVILY_API_KEY", "")
# 脚本模式同样固定 tavily provider（mock 的是 TavilyClient）
os.environ.setdefault("SEARCH_PROVIDER", "tavily")
# 脚本模式的数据目录隔离：conftest 不参与时，DATA_DIR 必须也指向临时目录，
# 否则下面的 clean_data_dirs() 会清空真实 backend/data（P0 事故防线）。
if not os.environ.get("DATA_DIR"):
    import tempfile
    os.environ["DATA_DIR"] = tempfile.mkdtemp(prefix="paperpilot-scripttest-data-")

BACKEND_DIR = Path(__file__).resolve().parent.parent
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

import unittest
import fitz
from docx import Document
from fastapi.testclient import TestClient

from app.config import settings
from app.main import app
from app.services.llm_client import LLMClient
from app.services.search_client import TavilyClient
from app.utils.ratelimit import RateLimitMiddleware


# ====== 常量 ======
API_KEY = "e2e-test-api-key-2026"
HEADERS = {"X-API-Key": API_KEY}
# P0 防线：测试读写的目录一律从 settings.data_dir 派生（conftest/上方 setdefault
# 已把它指向临时目录），绝不允许硬编码 BACKEND_DIR/"data"——那会指向真实论文库。
DATA_DIR = Path(settings.data_dir)
PAPERS_DIR = DATA_DIR / "papers"
PPTS_DIR = DATA_DIR / "ppts"
GITKEEP_NAME = ".gitkeep"
TEST_SAMPLES_DIR = "test_samples"  # 已提交进仓库的测试样本目录，清理时须跳过
# 真实数据目录（仅用于"拒绝清理"断言，绝不清它）
REAL_DATA_DIR = (BACKEND_DIR / "data").resolve()
REAL_PAPERS_DIR = REAL_DATA_DIR / "papers"
REAL_PPTS_DIR = REAL_DATA_DIR / "ppts"


# ====== Mock 数据 ======

MOCK_BREAKDOWN = {
    "summary": "本文提出 Transformer，完全基于注意力机制的新序列建模架构。",
    "background": "RNN/LSTM 存在并行性差的问题，限制训练效率。",
    "goal": "用纯注意力架构替代循环结构，提升并行度。",
    "method": "多头自注意力 + 位置编码 + 残差 + 前馈网络。",
    "experiment": "WMT 2014 英德翻译达到 28.4 BLEU，超越之前最佳。",
    "conclusion": "Transformer 在质量和并行性上均显著优于 RNN/CNN。",
    "key_points": ["完全基于注意力", "高度并行", "SOTA 性能"],
    "quotes": [{"section": "Abstract", "text": "Attention is all you need.", "page": 1}],
}

MOCK_INNOVATION = {
    "core_innovations": [
        {"title": "完全注意力架构", "description": "抛弃 RNN/CNN，仅用注意力。", "evidence": "[P3]", "page_ref": 3},
        {"title": "多头注意力机制", "description": "并行子空间学习。", "evidence": "[P4]", "page_ref": 4},
    ],
    "innovation_level": "disruptive",
    "level_reasoning": "开辟新范式，性能大幅领先。",
    "applicable_scenarios": ["机器翻译", "文本生成", "预训练"],
    "quotes": [],
}

MOCK_FLAWS = {
    "method_level": [{"description": "二次复杂度 O(n²)。", "evidence": "[P5]", "page_ref": 5, "severity": "major"}],
    "experiment_level": [{"description": "未在低资源语言上测试。", "evidence": "[P8]", "page_ref": 8, "severity": "minor"}],
    "writing_level": [],
    "improvements": [{"flaw_ref": 0, "suggestion": "引入稀疏注意力。", "feasibility": "high"}],
    "overall_assessment": "贡献突出但有优化空间。",
    "quotes": [],
}

MOCK_COMPARE = {
    "main_paper": {"title": "Attention Is All You Need", "year": 2017, "method_summary": "完全基于注意力的序列模型。"},
    "related_papers": [
        {"title": "BERT", "year": 2018, "url": "https://arxiv.org/abs/1810.04805", "method": "双向预训练", "dataset": "BookCorpus+Wiki", "result": "11 项 SOTA", "pros": "通用性强", "cons": "训练成本高", "relation_to_main": "改进"},
        {"title": "GPT-3", "year": 2020, "url": "https://arxiv.org/abs/2005.14165", "method": "自回归 LM", "dataset": "大规模网页", "result": "Few-shot SOTA", "pros": "少样本强", "cons": "推理成本极高", "relation_to_main": "互补"},
    ],
    "compare_table": [["维度", "主论文", "BERT", "GPT-3"], ["方法", "Transformer", "Encoder", "自回归 LM"], ["训练", "有监督", "自监督", "自回归"], ["关键结果", "28.4 BLEU", "11 SOTA", "Few-shot SOTA"]],
    "summary": "本文开创 Transformer，后续 BERT/GPT-3 基于其架构。",
    "main_advantages": ["开创性架构", "高效并行"],
    "main_disadvantages": ["长序列 O(n²)"],
}

MOCK_SEARCH_RESULTS = [
    {"title": "BERT", "url": "https://arxiv.org/abs/1810.04805", "content": "We introduce BERT.", "score": 0.95},
    {"title": "GPT-3", "url": "https://arxiv.org/abs/2005.14165", "content": "We train GPT-3.", "score": 0.92},
]


# ====== Sample 文件生成 ======

def create_sample_pdf() -> bytes:
    """生成多章节测试 PDF（4 页：Abstract+Intro / Method / Experiments / Conclusion）"""
    doc = fitz.open()
    pages = [
        "Attention Is All You Need\nAshish Vaswani et al.\nGoogle Brain 2017\n\nAbstract\nWe propose Transformer, based solely on attention.\n\nKeywords: attention, transformer\n\n1 Introduction\nRNNs suffer from sequential computation which prevents parallelization.\n",
        "2 Method\nThe Transformer uses stacked self-attention and fully connected layers.\n\n2.1 Scaled Dot-Product Attention\nInput consists of queries and keys of dimension dk.\n\n2.2 Multi-Head Attention\nLinearly project the queries, keys and values to different dimensions.\n",
        "3 Experiments\nOn WMT 2014 English-to-German, our model achieves 28.4 BLEU.\n\n3.1 Machine Translation\nOur big model outperforms all previously reported models.\n\n3.2 Model Variations\nWe varied the base model in different ways.\n",
        "4 Conclusion\nWe presented Transformer, the first sequence transduction model based entirely on attention.\n\nReferences\n[1] Bahdanau D et al. Neural machine translation. arXiv:1409.0473, 2014.\n",
    ]
    for text in pages:
        page = doc.new_page(width=595, height=842)
        y = 50
        for line in text.split("\n"):
            if line.strip():
                page.insert_text((50, y), line, fontsize=11)
            y += 15
            if y > 820:
                break
    out = doc.tobytes()
    doc.close()
    return out


def create_sample_docx() -> bytes:
    """生成多章节测试 DOCX"""
    doc = Document()
    doc.add_heading("Attention Is All You Need", level=0)
    doc.add_paragraph("Ashish Vaswani et al.")
    doc.add_heading("Abstract", level=1)
    doc.add_paragraph("We propose Transformer, based solely on attention mechanisms.")
    doc.add_heading("1 Introduction", level=1)
    doc.add_paragraph("RNNs suffer from sequential computation.")
    doc.add_heading("2 Method", level=1)
    doc.add_paragraph("The Transformer uses stacked self-attention.")
    doc.add_heading("2.1 Scaled Dot-Product Attention", level=2)
    doc.add_paragraph("Input consists of queries and keys.")
    for i in range(28):
        doc.add_paragraph(f"Filler paragraph {i+1} to span multiple pages.")
    doc.add_heading("3 Experiments", level=1)
    doc.add_paragraph("On WMT 2014, our model achieves 28.4 BLEU.")
    doc.add_heading("4 Conclusion", level=1)
    doc.add_paragraph("We presented Transformer, the first sequence transduction model based entirely on attention.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ====== 清理 / 限流 ======

def clean_data_dirs():
    # P0 双保险：任何情况下都拒绝清理真实数据目录。
    # 正常路径下 PAPERS_DIR/PPTS_DIR 派生自临时 DATA_DIR，永远不会命中这个分支；
    # 若未来有人把 DATA_DIR 改回真实目录，这里直接炸出来而不是静默删库。
    if (
        PAPERS_DIR.resolve() == REAL_PAPERS_DIR
        or PPTS_DIR.resolve() == REAL_PPTS_DIR
        or DATA_DIR.resolve() == REAL_DATA_DIR
    ):
        raise RuntimeError(
            "clean_data_dirs 拒绝执行：数据目录指向真实 backend/data！"
            f"（PAPERS_DIR={PAPERS_DIR}）请检查 conftest.py 的 DATA_DIR 注入。"
        )
    if PAPERS_DIR.exists():
        for entry in PAPERS_DIR.iterdir():
            if entry.name == GITKEEP_NAME:
                continue
            if entry.name == TEST_SAMPLES_DIR:
                continue
            if entry.is_dir():
                shutil.rmtree(entry, ignore_errors=True)
            else:
                entry.unlink(missing_ok=True)
    else:
        PAPERS_DIR.mkdir(parents=True, exist_ok=True)
        (PAPERS_DIR / GITKEEP_NAME).touch()
    if PPTS_DIR.exists():
        for entry in PPTS_DIR.iterdir():
            if entry.is_dir():
                shutil.rmtree(entry, ignore_errors=True)
            else:
                entry.unlink(missing_ok=True)
    else:
        PPTS_DIR.mkdir(parents=True, exist_ok=True)


def reset_rate_limit():
    RateLimitMiddleware._buckets.clear()


# ====== HTTP 辅助 ======

def upload_via_client(client, content: bytes, filename: str, headers: dict = None):
    if filename.endswith(".pdf"):
        media_type = "application/pdf"
    elif filename.endswith(".docx"):
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        media_type = "application/octet-stream"
    return client.post(
        "/api/upload",
        files={"file": (filename, io.BytesIO(content), media_type)},
        headers=headers or HEADERS,
    )


# ====== Mock 上下文 ======

class MockEnv:
    """统一管理 mock 上下文"""

    @staticmethod
    def _fake_build(*args, **kwargs):
        return [{"role": "system", "content": "fake system"}, {"role": "user", "content": "fake user"}]

    def __init__(self):
        self.stack = ExitStack()
        self.llm = None
        self.tavily = None
        from app.api import analyze as _analyze_mod
        from app.api import compare_papers as _compare_papers_mod
        self._analyze_mod = _analyze_mod
        self._compare_papers_mod = _compare_papers_mod

    def patch_llm(self, return_value):
        mock = AsyncMock(return_value=return_value)
        self.llm = self.stack.enter_context(patch.object(LLMClient, "chat_json", new=mock))
        return mock

    def patch_tavily(self, return_value):
        mock = AsyncMock(return_value=return_value)
        self.tavily = self.stack.enter_context(patch.object(TavilyClient, "search_related_papers", new=mock))
        return mock

    def patch_all_builds(self):
        for name in ("build_innovation_messages", "build_flaws_messages", "build_breakdown_messages",
                     "build_chapter_map_messages", "build_reduce_messages"):
            self.stack.enter_context(patch.object(self._analyze_mod, name, side_effect=self._fake_build))
        self.stack.enter_context(patch.object(self._compare_papers_mod, "build_compare_messages", side_effect=self._fake_build))
        from app.agent.prompts import compare as _compare_mod
        self.stack.enter_context(patch.object(_compare_mod, "build_compare_messages", side_effect=self._fake_build))
        return self

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        self.stack.close()


# ====== 共享 TestClient ======
_shared_client: TestClient = None


def get_client() -> TestClient:
    global _shared_client
    if _shared_client is None:
        _shared_client = TestClient(app)
    return _shared_client


# ====== E2E 测试基类（Sprint 4 / R2）======

class E2ETestBase(unittest.TestCase):
    """所有 E2E 测试的基类。

    提供：
    - setUpClass / tearDownClass：一次性环境校验
    - setUp / tearDown：每个测试前后清理 + mock 初始化
    - _upload_pdf / _upload_docx / _run_analyze：常用 helper
    """

    @classmethod
    def setUpClass(cls):
        clean_data_dirs()
        # 运行时覆盖限流值：e2e 生命周期请求量大，不能受 conftest 注入的 5/3 小值限制。
        # 限流中间件每次请求实时读 settings（见 app/utils/ratelimit.py:128-132），
        # 直接改属性即生效，无需重启进程。
        # 放在 setUpClass 而非模块级：pytest 在收集阶段会 import 全部 test 模块，
        # 模块级覆盖会让后执行的 test_auth_ratelimit 限流断言（5/3）失效，形成新竞态。
        settings.rate_limit_general_per_min = 10000
        settings.rate_limit_expensive_per_min = 10000
        assert settings.app_api_key == API_KEY, f"APP_API_KEY 未生效: {settings.app_api_key!r}"
        assert settings.auth_enabled is True, "鉴权应启用"
        assert settings.rate_limit_expensive_per_min == 10000, f"expensive 限流应为 10000，实际 {settings.rate_limit_expensive_per_min}"
        if not getattr(cls, "_config_printed", False):
            print("🧪 配置校验通过: api_key=***, general=10000/min, expensive=10000/min")
            cls._config_printed = True

    @classmethod
    def tearDownClass(cls):
        clean_data_dirs()

    def setUp(self):
        clean_data_dirs()
        reset_rate_limit()
        self.client = get_client()
        self.mocks = MockEnv()
        self.mocks.patch_all_builds()
        # 修复 Tavily client 单例 bug（源项目历史问题）
        import app.services.search_client as _sc
        if _sc._client is None:
            _sc._client = _sc.TavilyClient()

        def _fixed_get_tavily_client():
            if _sc._client is None:
                _sc._client = _sc.TavilyClient()
            return _sc._client

        self.mocks.stack.enter_context(patch.object(_sc, "get_tavily_client", _fixed_get_tavily_client))
        _sc._client.api_key = "tvly-fake-key-for-test"

    def tearDown(self):
        self.mocks.__exit__(None, None, None)
        clean_data_dirs()
        reset_rate_limit()

    # ---------- helper 方法（被子类复用） ----------

    def _upload_pdf(self) -> str:
        res = upload_via_client(self.client, create_sample_pdf(), "p.pdf")
        assert res.status_code == 200, f"上传 PDF 失败: {res.text}"
        return res.json()["data"]["paper_id"]

    def _upload_docx(self) -> str:
        res = upload_via_client(self.client, create_sample_docx(), "p.docx")
        assert res.status_code == 200, f"上传 DOCX 失败: {res.text}"
        return res.json()["data"]["paper_id"]

    def _run_analyze(self, paper_id, mock_data, type_):
        with self.mocks.patch_llm(mock_data):
            return self.client.post(
                "/api/analyze",
                json={"paper_id": paper_id, "type": type_},
                headers=HEADERS,
            )


# ====== 综合端到端流程（保留在主文件，因用到了所有 helper）======

class TestFullLifecycle(E2ETestBase):
    """完整生命周期：upload → analyze(3) → compare → ppt → download → delete"""

    def test_45_full_lifecycle(self):
        # 1. 上传
        paper_id = self._upload_pdf()
        self.assertTrue((PAPERS_DIR / paper_id).exists(), "上传后目录应存在")

        # 2. 列表能看到
        res = self.client.get("/api/papers", headers=HEADERS)
        self.assertEqual(res.json()["data"]["total"], 1)
        self.assertFalse(res.json()["data"]["items"][0]["has_breakdown"])

        # 3. 分析三种
        self._run_analyze(paper_id, MOCK_BREAKDOWN, "breakdown")
        self._run_analyze(paper_id, MOCK_INNOVATION, "innovation")
        self._run_analyze(paper_id, MOCK_FLAWS, "flaws")

        # 4. 列表里 has_* 应都 True
        item = self.client.get("/api/papers", headers=HEADERS).json()["data"]["items"][0]
        self.assertTrue(item["has_breakdown"])
        self.assertTrue(item["has_innovation"])
        self.assertTrue(item["has_flaws"])

        # 5. 读三种缓存
        for t in ("breakdown", "innovation", "flaws"):
            res = self.client.get(f"/api/analyze/{paper_id}/{t}", headers=HEADERS)
            self.assertEqual(res.status_code, 200, f"读 {t} 缓存失败: {res.text}")

        # 6. 联网对比
        with self.mocks.patch_tavily(MOCK_SEARCH_RESULTS), self.mocks.patch_llm(MOCK_COMPARE):
            res = self.client.post("/api/compare", json={"paper_id": paper_id, "max_results": 3}, headers=HEADERS)
        self.assertEqual(res.status_code, 200)

        # 7. 生成 PPT
        res = self.client.post(
            "/api/generate-ppt",
            json={"paper_id": paper_id, "include_flaws": True, "include_compare": True},
            headers=HEADERS,
        )
        self.assertEqual(res.status_code, 200)
        ppt_filename = res.json()["data"]["filename"]

        # 8-10. 下载 + 删除
        self.assertEqual(self.client.get(f"/api/download-ppt/{ppt_filename}", headers=HEADERS).status_code, 200)
        self.assertEqual(self.client.get(f"/api/papers/{paper_id}/file", headers=HEADERS).status_code, 200)
        self.assertEqual(self.client.delete(f"/api/papers/{paper_id}", headers=HEADERS).status_code, 200)
        self.assertFalse((PAPERS_DIR / paper_id).exists(), "删除后目录应不存在")
        print("✅ 完整生命周期通过（upload→analyze×3→compare→ppt→download×2→delete）")


# ====== 入口 ======

def main():
    """手动运行入口：按顺序跑所有 e2e 测试文件"""
    import glob
    import importlib
    print("\n" + "=" * 70)
    print("🚀 PaperPilot 端到端测试 (E2E) - 按职责拆分")
    print("=" * 70)

    pattern = str(Path(__file__).parent / "test_e2e*.py")
    files = sorted(glob.glob(pattern))
    print(f"📋 测试文件: {[Path(f).name for f in files]}\n")

    # 脚本模式不经过 conftest：tests/ 不在 sys.path，子文件里的裸名导入
    # `from test_e2e import ...` 会失败，这里补上（与 conftest.py 的注入一致）
    _tests_dir = str(Path(__file__).resolve().parent)
    if _tests_dir not in sys.path:
        sys.path.insert(0, _tests_dir)
    # 以裸名 "test_e2e" 导入基类模块：当前模块在脚本模式下叫 __main__ 或
    # tests.test_e2e，直接引用本模块的 E2ETestBase 会与子文件 `from test_e2e
    # import` 拿到的那份不是同一对象，issubclass 判断会落空（收集 0 个）。
    _base_mod = importlib.import_module("test_e2e")
    E2ETestBase = _base_mod.E2ETestBase

    loader = unittest.TestLoader()
    suite = unittest.TestSuite()
    for f in files:
        module_name = Path(f).stem
        mod = importlib.import_module(module_name)
        for name in dir(mod):
            obj = getattr(mod, name)
            if isinstance(obj, type) and issubclass(obj, unittest.TestCase) and obj is not unittest.TestCase:
                if issubclass(obj, E2ETestBase):
                    suite.addTests(loader.loadTestsFromTestCase(obj))

    test_names = [t._testMethodName for t in suite]
    print(f"📋 共 {len(test_names)} 个测试\n")

    passed = failed = 0
    errors = []
    start = time.time()
    for name in test_names:
        # 找对应的类
        cls = None
        for test in suite:
            if test._testMethodName == name:
                cls = type(test)
                break
        if cls is None:
            continue
        sub = unittest.TestSuite()
        sub.addTest(cls(name))
        result = unittest.TextTestRunner(verbosity=0, stream=open(os.devnull, "w")).run(sub)
        if result.wasSuccessful():
            passed += 1
        else:
            failed += 1
            for _t, err in result.failures + result.errors:
                errors.append((name, err))

    elapsed = time.time() - start
    print("\n" + "=" * 70)
    print(f"📊 结果: {passed} 通过 / {failed} 失败 / 共 {passed + failed}（耗时 {elapsed:.1f}s）")
    if errors:
        print("\n❌ 失败详情：\n")
        for name, err in errors:
            print(f"--- {name} ---")
            print(err)
    else:
        print("🎉 全部通过！")
    print("=" * 70)
    return failed == 0


if __name__ == "__main__":
    sys.exit(0 if main() else 1)
