"""
DOCX 解析器测试（不依赖 LLM，pytest 下可独立/乱序运行）

- 样本用 tmp_path 构造，不写入仓库 data/papers/test_samples/（避免污染真实样本）
- 每个用例独立拿到样本路径，无内部顺序依赖
- test_with_real_docx 用仓库自带的 transformer_sample.docx，不存在则 skip

用法：
    cd backend
    python -m pytest tests/test_docx_parser.py -v
"""
import sys
from pathlib import Path

import pytest

BACKEND_DIR = Path(__file__).resolve().parent.parent
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from app.services.docx_parser import DOCXParser, parse_docx

TEST_SAMPLE_DOCX = BACKEND_DIR / "data" / "papers" / "test_samples" / "transformer_sample.docx"


def create_sample_docx(target: Path) -> Path:
    """用 python-docx 构造示例 DOCX 写入 target（临时目录，不污染仓库）"""
    try:
        from docx import Document
    except ImportError:
        print("❌ 需要 python-docx: pip install python-docx")
        sys.exit(1)

    doc = Document()

    # 标题
    doc.add_heading("Attention Is All You Need", level=1)

    # 作者
    doc.add_paragraph("Ashish Vaswani, Noam Shazeer, Niki Parmar, Jakob Uszkoreit, Llion Jones")

    # 摘要
    doc.add_paragraph("Abstract")
    doc.add_paragraph(
        "We propose a new simple network architecture, the Transformer, "
        "based solely on attention mechanisms, dispensing with recurrence "
        "and convolutions entirely. Experiments on two machine translation "
        "tasks show these models to be superior in quality while being more "
        "parallelizable and requiring significantly less time to train."
    )

    # 关键词
    doc.add_paragraph("Keywords: attention, transformer, neural networks")

    # 章节 1
    doc.add_heading("1 Introduction", level=1)
    doc.add_paragraph(
        "Recurrent neural networks, long short-term memory and gated recurrent "
        "neural networks have been firmly established as state-of-the-art approaches "
        "in sequence modeling and transduction problems."
    )
    doc.add_paragraph(
        "However, RNNs suffer from sequential computation which prevents parallelization "
        "within training examples."
    )

    # 章节 2
    doc.add_heading("2 Background", level=1)
    doc.add_paragraph(
        "The goal of reducing sequential computation has been addressed by various models. "
        "Extended Neural GPU and ConvS2S use convolutional neural networks as building blocks."
    )

    # 章节 3
    doc.add_heading("3 Method", level=1)
    doc.add_heading("3.1 Scaled Dot-Product Attention", level=2)
    doc.add_paragraph(
        "We call our particular attention Scaled Dot-Product Attention. "
        "The input consists of queries and keys of dimension dk."
    )

    # 多塞点段落让 pages 切分触发
    for i in range(30):
        doc.add_paragraph(f"This is filler paragraph {i+1} to ensure the document spans multiple pages.")

    # 章节 4
    doc.add_heading("4 Experiments", level=1)
    doc.add_paragraph(
        "We evaluate our model on two machine translation tasks. On WMT 2014 "
        "English-to-German translation task, our model achieves 28.4 BLEU."
    )

    # 章节 5
    doc.add_heading("5 Conclusion", level=1)
    doc.add_paragraph(
        "In this work, we presented the Transformer, the first sequence transduction "
        "model based entirely on attention."
    )

    # 保存到临时目录
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(target))
    return target


@pytest.fixture(scope="module")
def sample_docx(tmp_path_factory) -> Path:
    """模块级构造示例 DOCX 到临时目录，消除 test_metadata 与其他用例的顺序依赖"""
    target = tmp_path_factory.mktemp("docx_samples") / "transformer_sample.docx"
    return create_sample_docx(target)


def test_metadata(sample_docx: Path):
    """测试元数据提取"""
    print("=" * 60)
    print("🧪 测试元数据提取")
    print("=" * 60)
    print(f"样本: {sample_docx}\n")

    with DOCXParser(sample_docx) as parser:
        meta = parser.extract_metadata()
        print(f"📌 标题: {meta.title}")
        print(f"👥 作者: {meta.authors}")
        print(f"📅 年份: {meta.year}")
        print(f"🔑 关键词: {meta.keywords}")
        print(f"📝 摘要: {meta.abstract[:150]}...")
        print()

        # 断言
        assert "Attention" in meta.title, f"标题应该含 'Attention'，实际 '{meta.title}'"
        assert len(meta.authors) >= 2, f"作者数应该 >=2，实际 {len(meta.authors)}"
        assert "Transformer" in meta.abstract, "摘要应该含 'Transformer'"
        assert len(meta.keywords) >= 1, f"关键词应该 >=1，实际 {len(meta.keywords)}"

    print("✅ 通过\n")


def test_pages(sample_docx: Path):
    """测试按段分块"""
    print("=" * 60)
    print("🧪 测试分页（按段分块）")
    print("=" * 60)

    with DOCXParser(sample_docx) as parser:
        pages = parser.extract_pages()
        print(f"\n共 {len(pages)} 个伪页（每 25 段一页）")
        for p in pages[:3]:
            print(f"--- 第 {p.page_num} 页 ({len(p.text)} 字符) ---")
            print(p.text[:100].replace("\n", " ") + "...")

        assert len(pages) >= 2, f"应该 >=2 页，实际 {len(pages)}"
        for p in pages:
            assert p.page_num >= 1
            assert p.text or p.page_num == 1  # 第一页可能为空

    print("✅ 通过\n")


def test_toc(sample_docx: Path):
    """测试 TOC 提取（基于 Heading 1）"""
    print("=" * 60)
    print("🧪 测试 TOC 提取")
    print("=" * 60)

    with DOCXParser(sample_docx) as parser:
        toc = parser.extract_toc()
        print(f"\n提取到 {len(toc)} 条大纲：")
        for level, title, page in toc:
            print(f"  [{level}] {title} (P{page})")

        assert len(toc) >= 3, f"应该 >=3 条大纲，实际 {len(toc)}"

        # 验证能找到 Introduction / Method / Experiment / Conclusion
        titles = [t[1].lower() for t in toc]
        assert any("introduction" in t for t in titles), "应该找到 Introduction"
        assert any("method" in t for t in titles), "应该找到 Method"
        assert any("experiment" in t for t in titles), "应该找到 Experiment"
        assert any("conclusion" in t for t in titles), "应该找到 Conclusion"

    print("✅ 通过\n")


def test_full_parse(sample_docx: Path):
    """测试一站式 parse()"""
    print("=" * 60)
    print("🧪 测试完整 parse()")
    print("=" * 60)

    parsed = parse_docx(sample_docx)

    # 验证输出结构与 PDFParser 一致
    assert "meta" in parsed
    assert "page_count" in parsed
    assert "full_text" in parsed
    assert "toc" in parsed
    assert "pages" in parsed

    print(f"\nmeta.title = {parsed['meta']['title']}")
    print(f"page_count = {parsed['page_count']}")
    print(f"toc 条数 = {len(parsed['toc'])}")
    print(f"full_text 长度 = {len(parsed['full_text'])} 字符")
    print(f"pages 数量 = {len(parsed['pages'])}")

    # 验证 page 格式
    for p in parsed["pages"][:2]:
        assert "page_num" in p
        assert "text" in p
        assert "width" in p
        assert "height" in p
        assert "blocks" in p

    # 页码标记协议：full_text 必须带 === Page N === 标记（legacy 分析依赖）
    import re
    markers = re.findall(r"=== Page (\d+) ===", parsed["full_text"])
    assert markers, "full_text 应包含 === Page N === 页码标记"
    assert int(markers[0]) == 1, "首个标记页码应为 1"
    assert len(markers) == parsed["page_count"], (
        f"标记数 {len(markers)} 应与伪页数 {parsed['page_count']} 一致"
    )

    print("✅ 通过\n")


def test_abstract_long_paragraph(tmp_path):
    """测试（B11 回归）：Abstract 段落超过 100 字时摘要不丢失"""
    from docx import Document

    doc = Document()
    doc.add_heading("A Study of Long Abstracts", level=1)
    doc.add_paragraph("Alice Wonder, Bob Builder, Carol Chan")
    # 单段长摘要（>100 字），带 "Abstract:" 前缀
    long_abstract = "Abstract: " + (
        "This work studies a very long abstract paragraph that exceeds "
        "one hundred characters in total length. " * 3
    )
    doc.add_paragraph(long_abstract)
    doc.add_heading("1 Introduction", level=1)
    doc.add_paragraph("Introductory sentence for the introduction section.")

    path = tmp_path / "long_abstract.docx"
    doc.save(str(path))

    parsed = parse_docx(path)
    abstract = parsed["meta"]["abstract"]
    print(f"\n提取到摘要: {abstract[:120]}...")
    assert "long abstract paragraph" in abstract, (
        f"长段落摘要应被完整提取，实际: {abstract!r}"
    )


def test_with_real_docx():
    """测试仓库自带的真实 DOCX（不存在则跳过）"""
    print("=" * 60)
    print("🧪 测试真实 DOCX（可选）")
    print("=" * 60)

    if not TEST_SAMPLE_DOCX.exists():
        print(f"⚠️  样本不存在: {TEST_SAMPLE_DOCX}，跳过\n")
        pytest.skip(f"样本不存在: {TEST_SAMPLE_DOCX}")

    docx_path = TEST_SAMPLE_DOCX
    parsed = parse_docx(docx_path)
    print(f"\n📄 {docx_path.name}")
    print(f"标题: {parsed['meta']['title']}")
    print(f"作者: {parsed['meta']['authors']}")
    print(f"页数: {parsed['page_count']} (伪页)")
    print(f"章节: {len(parsed['toc'])} 条")
    print(f"全文: {len(parsed['full_text'])} 字符")
    print(f"摘要: {parsed['meta']['abstract'][:200]}")
    # 真实断言：输出结构完整、标题非空
    assert parsed["page_count"] >= 1, "真实 DOCX 应至少解析出 1 页"
    assert parsed["meta"]["title"], "真实 DOCX 应能提取到标题"
    print("✅ 完成\n")


def main():
    """手动运行入口：直接以 pytest 运行本文件（保证与 CI 行为一致）"""
    import pytest as _pytest
    sys.exit(_pytest.main([str(Path(__file__).resolve())]))


if __name__ == "__main__":
    main()
